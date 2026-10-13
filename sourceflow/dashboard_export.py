"""Dashboard export: the list, the summary, and every chart on screen.

The browser sends what it is showing - the case rows, the KPI summary, and
for each chart its title, its data table and a PNG of the canvas. This turns
that into one workbook: a Cases sheet, a Summary sheet, and a sheet per chart
holding the data table, a native Excel chart built from it (so it can be
re-styled or re-pointed at other data), and the picture of what the buyer
actually saw. openpyxl draws the native charts; Pillow lets it embed the PNG.
"""
import base64
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, RadarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD_FILL = PatternFill("solid", fgColor="E7F3EB")
_HEAD_FONT = Font(bold=True, color="00504A")
_TITLE_FONT = Font(bold=True, size=13)
_MUTED_FONT = Font(color="6E8290", size=9)

# Chart.js type -> the Excel chart that says the same thing. Anything not
# listed (bubble, scatter) gets its data and picture but no native chart.
_NATIVE = {
    "bar": "bar", "horizontalBar": "bar", "line": "line",
    "doughnut": "doughnut", "pie": "pie", "polarArea": "bar", "radar": "radar",
    "stacked": "stacked", "table": "stacked", "value": "bar",
}


def _sheet_name(title: str, taken: set) -> str:
    base = re.sub(r"\s+", " ", re.sub(r"[\[\]:*?/\\]", " ", title or "Chart")).strip()[:28].strip() or "Chart"
    name, n = base, 2
    while name.lower() in taken:
        name = f"{base[:25]} {n}"; n += 1
    taken.add(name.lower())
    return name


def _cell(v):
    """Anything Excel cannot hold as a value becomes text rather than an
    error - the browser sends what it has, and a stray list must not sink
    the whole workbook."""
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def _write_rows(ws, rows: list, start_row: int = 1) -> tuple:
    """A list of dicts as a header row plus data rows. Returns (first, last)."""
    if not rows:
        ws.cell(row=start_row, column=1, value="Nothing in this selection").font = _MUTED_FONT
        return start_row, start_row
    headers = list(rows[0].keys())
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h); c.fill = _HEAD_FILL; c.font = _HEAD_FONT
    for i, r in enumerate(rows, start_row + 1):
        for j, h in enumerate(headers, 1):
            ws.cell(row=i, column=j, value=_cell(r.get(h)))
    for j, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r.get(h) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 48)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row, start_row + len(rows)


def _native_chart(kind: str, ws, first_row: int, last_row: int, n_series: int, title: str, stacked: bool):
    """Build the Excel chart over the data table just written: labels in
    column A, one series per further column, header row on top."""
    if kind in ("bar", "stacked"):
        ch = BarChart(); ch.type = "col"
        if stacked or kind == "stacked":
            ch.grouping = "stacked"; ch.overlap = 100
    elif kind == "line":
        ch = LineChart()
    elif kind == "doughnut":
        ch = DoughnutChart(); ch.holeSize = 55
    elif kind == "pie":
        ch = PieChart()
    elif kind == "radar":
        ch = RadarChart(); ch.type = "standard"
    else:
        return None
    data = Reference(ws, min_col=2, max_col=1 + n_series, min_row=first_row, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=first_row + 1, max_row=last_row)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.title = title
    ch.height, ch.width = 9, 16
    if kind in ("doughnut", "pie") and n_series == 1:
        ch.legend.position = "r"
    return ch


def _png(data_url: str):
    if not data_url or "," not in data_url:
        return None
    try:
        raw = base64.b64decode(data_url.split(",", 1)[1])
        img = XLImage(io.BytesIO(raw))
    except Exception:  # noqa: BLE001 - a bad picture never breaks the export
        return None
    # Fit a dashboard canvas into the sheet at a readable size.
    if img.width and img.width > 720:
        ratio = 720 / img.width
        img.width, img.height = 720, int(img.height * ratio)
    return img


def build(payload: dict) -> bytes:
    wb = Workbook()
    taken = set()
    first = True
    for sheet in payload.get("sheets") or []:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = _sheet_name(sheet.get("name") or "Sheet", taken)
        _write_rows(ws, sheet.get("rows") or [])

    for chart in payload.get("charts") or []:
        title = (chart.get("title") or "Chart").strip()
        ws = wb.create_sheet(_sheet_name(title, taken))
        ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
        if chart.get("subtitle"):
            ws.cell(row=2, column=1, value=chart["subtitle"]).font = _MUTED_FONT

        headers = chart.get("headers") or []
        rows = chart.get("rows") or []
        table_row = 4
        for j, h in enumerate(headers, 1):
            c = ws.cell(row=table_row, column=j, value=h); c.fill = _HEAD_FILL; c.font = _HEAD_FONT
        for i, r in enumerate(rows, table_row + 1):
            for j, v in enumerate(r, 1):
                ws.cell(row=i, column=j, value=_cell(v))
        for j, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[j - 1] if j - 1 < len(r) else "")) for r in rows[:100]])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 40)
        ws.cell(row=table_row, column=1).alignment = Alignment(horizontal="left")

        # How many columns after the label are series to chart; a table may
        # carry totals or percentages after them that belong in the sheet
        # but not on the chart.
        n_series = max(0, min(int(chart.get("series") or len(headers) - 1), len(headers) - 1))
        anchor_col = get_column_letter(len(headers) + 2) if headers else "D"
        kind = _NATIVE.get(chart.get("type") or "")
        placed = False
        if kind and rows and n_series >= 1:
            native = _native_chart(kind, ws, table_row, table_row + len(rows), n_series, title, bool(chart.get("stacked")))
            if native is not None:
                ws.add_chart(native, f"{anchor_col}{table_row}")
                placed = True
        img = _png(chart.get("png"))
        if img is not None:
            # Below the native chart if there is one, else where it would be.
            row = table_row + 20 if placed else table_row
            ws.cell(row=row - 1, column=len(headers) + 2 if headers else 4, value="As shown on the dashboard").font = _MUTED_FONT
            ws.add_image(img, f"{anchor_col}{row}")

    if len(wb.worksheets) == 0:
        wb.active.title = "Empty"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
